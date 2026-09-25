"""Item Checker - desktop app.

Stores the item list in a SQLite database (items.db, next to this file),
compares Physical Stock against Committed Stock, and flags items that need attention.
Import/export: CSV, Excel (.xlsx), and SQLite (.db) backup/restore.
"""
import csv
import datetime
import os
import re
import sqlite3
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    openpyxl = None

if getattr(sys, 'frozen', False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get('ITEMCHECKER_DB') or os.path.join(APP_DIR, 'items.db')

HEADERS = ['Item Name', 'SKU', 'Physical Stock QTY', 'Committed Stock', 'Active']
STATUSES = ['OVER-COMMITTED', 'NO STOCK', 'LOW', 'OK']

SCHEMA = """
CREATE TABLE IF NOT EXISTS items (
  id              INTEGER PRIMARY KEY AUTOINCREMENT,
  item_name       TEXT    NOT NULL,
  sku             TEXT    NOT NULL UNIQUE COLLATE NOCASE,
  physical_stock  INTEGER NOT NULL DEFAULT 0 CHECK (physical_stock  >= 0),
  committed_stock INTEGER NOT NULL DEFAULT 0 CHECK (committed_stock >= 0),
  active          INTEGER NOT NULL DEFAULT 1 CHECK (active IN (0,1)),
  created_at      TEXT    NOT NULL DEFAULT (datetime('now','localtime')),
  updated_at      TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
);
CREATE TRIGGER IF NOT EXISTS trg_items_updated AFTER UPDATE ON items
BEGIN
  UPDATE items SET updated_at = datetime('now','localtime') WHERE id = OLD.id;
END;
CREATE VIEW IF NOT EXISTS v_item_check AS
SELECT id, item_name, sku, physical_stock, committed_stock,
       (physical_stock - committed_stock) AS available_stock, active,
       CASE WHEN committed_stock > physical_stock       THEN 'OVER-COMMITTED'
            WHEN physical_stock = 0                     THEN 'NO STOCK'
            WHEN physical_stock - committed_stock <= 5  THEN 'LOW'
            ELSE 'OK' END AS status,
       updated_at
FROM items;
"""

SAMPLE_ITEMS = [
    ('Nordic Dining Chair', 'HC-CHR-001', 120, 30, 1),
    ('Oak Coffee Table', 'HC-TBL-014', 15, 18, 1),
    ('Rattan Floor Lamp', 'HC-LMP-077', 0, 0, 0),
    ('Linen Throw Pillow', 'HC-PLW-203', 40, 36, 1),
    ('Ceramic Vase Set', 'HC-DEC-310', 60, 12, 1),
]

PRESETS = [
    ('Summary by status', """SELECT status, COUNT(*) AS items,
       SUM(physical_stock) AS physical, SUM(committed_stock) AS committed,
       SUM(available_stock) AS available
FROM v_item_check
GROUP BY status
ORDER BY CASE status WHEN 'OVER-COMMITTED' THEN 1 WHEN 'NO STOCK' THEN 2 WHEN 'LOW' THEN 3 ELSE 4 END;"""),
    ('Over-committed', """SELECT item_name, sku, physical_stock, committed_stock, available_stock
FROM v_item_check
WHERE status = 'OVER-COMMITTED'
ORDER BY available_stock;"""),
    ('Low / no stock (active)', """SELECT item_name, sku, physical_stock, committed_stock, available_stock, status
FROM v_item_check
WHERE active = 1 AND status IN ('LOW', 'NO STOCK')
ORDER BY available_stock;"""),
    ('Inactive items', """SELECT item_name, sku, physical_stock, committed_stock
FROM items
WHERE active = 0
ORDER BY item_name;"""),
    ('Recently updated', """SELECT item_name, sku, physical_stock, committed_stock, updated_at
FROM items
ORDER BY updated_at DESC
LIMIT 20;"""),
    ('Schema', "SELECT type, name, sql FROM sqlite_master WHERE name NOT LIKE 'sqlite_%';"),
]


def compute_status(physical, committed):
    if committed > physical:
        return 'OVER-COMMITTED'
    if physical == 0:
        return 'NO STOCK'
    if physical - committed <= 5:
        return 'LOW'
    return 'OK'


# ---------------------------------------------------------------- Import parsing

HEADER_ALIASES = {
    'item_name': ['itemname', 'name', 'item', 'productname'],
    'sku': ['sku', 'itemcode', 'skucode'],
    'physical_stock': ['physicalstockqty', 'physicalstock', 'physicalqty', 'physical', 'stockonhand', 'onhand', 'qtyonhand'],
    'committed_stock': ['committedstock', 'committedstockqty', 'committedqty', 'committed'],
    'active': ['active', 'isactive'],
}


def norm_header(h):
    return re.sub(r'[^a-z0-9]', '', str(h if h is not None else '').lower())


def map_headers(row):
    found = {}
    for i, h in enumerate(row):
        n = norm_header(h)
        for field, aliases in HEADER_ALIASES.items():
            if field not in found and n in aliases:
                found[field] = i
    return found


def cell_text(v):
    """Cell value as text, the way it would look in Excel."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_qty(v):
    """Whole number >= anything; None if not a whole number. Blank = 0."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v.is_integer() else None
    s = re.sub(r'[,\s]', '', cell_text(v))
    if s == '':
        return 0
    if not re.fullmatch(r'-?\d+(\.0+)?', s):
        return None
    return int(s.split('.')[0])


def parse_active(v):
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)) and v in (0, 1):
        return int(v)
    s = cell_text(v).lower()
    if s == '':
        return 1
    if s in ('true', 'yes', 'y', '1', 'active'):
        return 1
    if s in ('false', 'no', 'n', '0', 'inactive'):
        return 0
    return None


def read_table_file(path):
    """Return the first sheet of a .csv or .xlsx file as a list of rows."""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        with open(path, 'rb') as f:
            raw = f.read()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('cp1252', errors='replace')
        return list(csv.reader(text.splitlines()))
    if ext in ('.xlsx', '.xlsm'):
        if openpyxl is None:
            raise RuntimeError('Excel support needs openpyxl.  Run:  pip install openpyxl')
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            return [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    raise RuntimeError(f'Unsupported file type "{ext}". Use .csv, .xlsx or a .db backup.')


def split_sql(sql):
    """Split a script into single statements (respecting quotes/triggers)."""
    parts, buf = [], ''
    for chunk in re.split(r'(;)', sql):
        buf += chunk
        if chunk == ';' and sqlite3.complete_statement(buf):
            parts.append(buf)
            buf = ''
    if buf.strip():
        parts.append(buf)
    return [p for p in parts if p.strip().strip(';').strip()]


# ---------------------------------------------------------------- Data layer

class Store:
    def __init__(self, path):
        self.path = path
        is_new = not os.path.exists(path)
        self.conn = sqlite3.connect(path, isolation_level=None)
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript(SCHEMA)
        if is_new:
            self.conn.execute('BEGIN')
            self.conn.executemany(
                'INSERT INTO items (item_name, sku, physical_stock, committed_stock, active) VALUES (?,?,?,?,?)',
                SAMPLE_ITEMS)
            self.conn.execute('COMMIT')
            self.sample_flag = True

    # PRAGMA user_version = 1 marks "sample data still loaded" (shows the banner).
    @property
    def sample_flag(self):
        return self.conn.execute('PRAGMA user_version').fetchone()[0] == 1

    @sample_flag.setter
    def sample_flag(self, on):
        self.conn.execute(f'PRAGMA user_version = {1 if on else 0}')

    def close(self):
        self.conn.close()

    def list_items(self, search='', active='all', status='all', sort_col='item_name', sort_dir='asc'):
        where, params = [], []
        if search:
            like = '%' + re.sub(r'([\\%_])', r'\\\1', search) + '%'
            where.append("(item_name LIKE ? ESCAPE '\\' OR sku LIKE ? ESCAPE '\\')")
            params += [like, like]
        if active in ('0', '1'):
            where.append('active = ?')
            params.append(int(active))
        if status in STATUSES:
            where.append('status = ?')
            params.append(status)
        cols = ['item_name', 'sku', 'physical_stock', 'committed_stock', 'available_stock', 'status', 'active', 'updated_at']
        col = sort_col if sort_col in cols else 'item_name'
        collate = ' COLLATE NOCASE' if col in ('item_name', 'sku', 'status') else ''
        direction = 'DESC' if sort_dir == 'desc' else 'ASC'
        sql = (f"SELECT * FROM v_item_check {'WHERE ' + ' AND '.join(where) if where else ''} "
               f"ORDER BY {col}{collate} {direction}, item_name COLLATE NOCASE")
        return self.conn.execute(sql, params).fetchall()

    def totals(self):
        return self.conn.execute("""
            SELECT COUNT(*) AS total,
                   COALESCE(SUM(active), 0) AS active,
                   COALESCE(SUM(physical_stock), 0) AS physical,
                   COALESCE(SUM(committed_stock), 0) AS committed,
                   COALESCE(SUM(status = 'OVER-COMMITTED'), 0) AS over
            FROM v_item_check""").fetchone()

    def get(self, item_id):
        return self.conn.execute('SELECT * FROM items WHERE id = ?', (item_id,)).fetchone()

    def find_sku(self, sku, except_id=None):
        return self.conn.execute('SELECT id, item_name, sku FROM items WHERE sku = ? AND id IS NOT ?',
                                 (sku, except_id)).fetchone()

    def save_item(self, item_id, name, sku, physical, committed, active):
        """Validate and insert (item_id None) or update. Raises ValueError with a user message."""
        name, sku = (name or '').strip(), (sku or '').strip()
        if not name:
            raise ValueError('Item Name is required.')
        if not sku:
            raise ValueError('SKU is required.')
        for label, v in (('Physical Stock QTY', physical), ('Committed Stock', committed)):
            if v is None:
                raise ValueError(f'{label} must be a whole number.')
            if v < 0:
                raise ValueError(f'{label} cannot be negative.')
        dup = self.find_sku(sku, item_id)
        if dup:
            raise ValueError(f'SKU "{dup["sku"]}" already exists ({dup["item_name"]}). SKUs are not case-sensitive.')
        if item_id is None:
            self.conn.execute('INSERT INTO items (item_name, sku, physical_stock, committed_stock, active) VALUES (?,?,?,?,?)',
                              (name, sku, physical, committed, int(bool(active))))
        else:
            self.conn.execute('UPDATE items SET item_name=?, sku=?, physical_stock=?, committed_stock=?, active=? WHERE id=?',
                              (name, sku, physical, committed, int(bool(active)), item_id))

    def toggle_active(self, item_id):
        self.conn.execute('UPDATE items SET active = 1 - active WHERE id = ?', (item_id,))

    def delete(self, item_id):
        self.conn.execute('DELETE FROM items WHERE id = ?', (item_id,))

    def clear_all(self):
        self.conn.execute('BEGIN')
        self.conn.execute('DELETE FROM items')
        self.conn.execute("DELETE FROM sqlite_sequence WHERE name = 'items'")
        self.conn.execute('COMMIT')
        self.sample_flag = False

    def import_rows(self, rows):
        """UPSERT rows (first sheet incl. header). Returns (added, updated, skipped[(row_no, sku, reason)])."""
        header_idx, cols = -1, None
        for i, row in enumerate(rows[:10]):
            m = map_headers(row)
            if 'item_name' in m and 'sku' in m:
                header_idx, cols = i, m
                break
        if header_idx < 0:
            raise ValueError('Could not find the header row. The file needs at least the columns "Item Name" and "SKU".')

        def get(row, field):
            i = cols.get(field)
            return row[i] if i is not None and i < len(row) else None

        added = updated = 0
        skipped = []
        self.conn.execute('BEGIN')
        try:
            for i in range(header_idx + 1, len(rows)):
                row, row_no = rows[i], i + 1
                if not row or all(cell_text(c) == '' for c in row):
                    continue
                name, sku = cell_text(get(row, 'item_name')), cell_text(get(row, 'sku'))
                phys, comm = parse_qty(get(row, 'physical_stock')), parse_qty(get(row, 'committed_stock'))
                active = parse_active(get(row, 'active'))
                why = None
                if not name:
                    why = 'Item Name is required'
                elif not sku:
                    why = 'SKU is required'
                elif phys is None:
                    why = f'Physical Stock QTY "{cell_text(get(row, "physical_stock"))}" is not a whole number'
                elif comm is None:
                    why = f'Committed Stock "{cell_text(get(row, "committed_stock"))}" is not a whole number'
                elif phys < 0:
                    why = 'Physical Stock QTY cannot be negative'
                elif comm < 0:
                    why = 'Committed Stock cannot be negative'
                elif active is None:
                    why = f'Active "{cell_text(get(row, "active"))}" is not TRUE/FALSE, YES/NO, 1/0 or Active'
                if why:
                    skipped.append((row_no, sku, why))
                    continue
                exists = self.find_sku(sku) is not None
                self.conn.execute("""
                    INSERT INTO items (item_name, sku, physical_stock, committed_stock, active) VALUES (?,?,?,?,?)
                    ON CONFLICT(sku) DO UPDATE SET item_name = excluded.item_name,
                      physical_stock = excluded.physical_stock, committed_stock = excluded.committed_stock,
                      active = excluded.active""", (name, sku, phys, comm, active))
                if exists:
                    updated += 1
                else:
                    added += 1
            self.conn.execute('COMMIT')
        except Exception:
            self.conn.execute('ROLLBACK')
            raise
        return added, updated, skipped

    def export_rows(self):
        return self.conn.execute('SELECT item_name, sku, physical_stock, committed_stock, active FROM items '
                                 'ORDER BY item_name COLLATE NOCASE, sku').fetchall()

    def backup_to(self, path):
        if os.path.abspath(path) == os.path.abspath(self.path):
            raise ValueError('Choose a different file name than the live database.')
        if os.path.exists(path):
            os.remove(path)
        dst = sqlite3.connect(path)
        try:
            self.conn.backup(dst)
        finally:
            dst.close()

    def restore_from(self, path):
        """Replace all data with a .db backup. The backup file itself is not modified."""
        src = sqlite3.connect(path)
        mem = sqlite3.connect(':memory:')
        try:
            try:
                src.backup(mem)
                if not mem.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='items'").fetchone():
                    raise ValueError('This .db file does not contain an items table.')
                mem.executescript(SCHEMA)
                mem.execute('SELECT id, item_name, sku, status FROM v_item_check LIMIT 1').fetchall()
            except sqlite3.DatabaseError as ex:
                raise ValueError(f'This file is not a valid Item Checker database ({ex}).')
            mem.execute('PRAGMA user_version = 0')
            mem.backup(self.conn)
        finally:
            src.close()
            mem.close()
        return self.conn.execute('SELECT COUNT(*) FROM items').fetchone()[0]

    def run_sql(self, sql):
        """Run one or more statements. Returns (columns, rows, rows_changed, any_change)."""
        before = self.conn.total_changes
        columns, rows, changed = None, None, 0
        for stmt in split_sql(sql):
            cur = self.conn.execute(stmt)
            if cur.description:
                columns = [d[0] for d in cur.description]
                rows = [tuple(r) for r in cur.fetchall()]
            elif cur.rowcount > 0:
                changed += cur.rowcount
        return columns, rows, changed, self.conn.total_changes > before


# ---------------------------------------------------------------- File writers

def style_header(ws):
    fill = PatternFill('solid', fgColor='0F766E')
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = fill
    for col, w in zip('ABCDE', (34, 16, 18, 16, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'


def write_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Items'
    ws.append(HEADERS)
    for r in rows:
        ws.append([r['item_name'], r['sku'], r['physical_stock'], r['committed_stock'], bool(r['active'])])
    style_header(ws)
    wb.save(path)


def write_csv(path, header, rows):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Items'
    ws.append(HEADERS)
    style_header(ws)
    help_ws = wb.create_sheet('Instructions')
    for line in [
        ['How to fill the Items sheet'],
        ['Item Name and SKU are required.'],
        ['SKU must be unique. It is not case-sensitive (HC-001 = hc-001). An existing SKU is updated on import.'],
        ['Physical Stock QTY and Committed Stock: whole numbers, 0 or more. Blank = 0. Commas are fine (1,200).'],
        ['Active: TRUE/FALSE, YES/NO, 1/0, or Active/Inactive. Blank = TRUE.'],
        [],
        ['Example'],
        HEADERS,
        ['Nordic Dining Chair', 'HC-CHR-001', 120, 30, 'TRUE'],
        ['Oak Coffee Table', 'HC-TBL-014', 15, 18, 'TRUE'],
        ['Rattan Floor Lamp', 'HC-LMP-077', 0, 0, 'FALSE'],
    ]:
        help_ws.append(line)
    help_ws['A1'].font = Font(bold=True, size=12)
    help_ws.column_dimensions['A'].width = 34
    wb.save(path)


# ---------------------------------------------------------------- UI

TEAL, TEAL_DARK, BG, CARD, MUTED, RED = '#0f766e', '#115e59', '#f1f5f9', '#ffffff', '#64748b', '#dc2626'
ROW_COLORS = {'OVER-COMMITTED': '#fee2e2', 'NO STOCK': '#e2e8f0', 'LOW': '#fef3c7', 'OK': '#ffffff'}
FONT = ('Segoe UI', 10)


def today():
    return datetime.date.today().isoformat()


class ItemDialog(tk.Toplevel):
    def __init__(self, app, item=None):
        super().__init__(app)
        self.app, self.item = app, item
        self.title('Edit Item' if item else 'Add Item')
        self.configure(bg=CARD, padx=18, pady=14)
        self.resizable(False, False)
        self.transient(app)

        self.v_name = tk.StringVar(value=item['item_name'] if item else '')
        self.v_sku = tk.StringVar(value=item['sku'] if item else '')
        self.v_phys = tk.StringVar(value=str(item['physical_stock']) if item else '0')
        self.v_comm = tk.StringVar(value=str(item['committed_stock']) if item else '0')
        self.v_active = tk.BooleanVar(value=bool(item['active']) if item else True)

        def label(text, r, c=0):
            tk.Label(self, text=text, bg=CARD, font=('Segoe UI Semibold', 10)).grid(row=r, column=c, sticky='w', pady=(6, 2))

        label('Item Name *', 0)
        e_name = ttk.Entry(self, textvariable=self.v_name, width=44, font=FONT)
        e_name.grid(row=1, column=0, columnspan=2, sticky='we')
        label('SKU *', 2)
        ttk.Entry(self, textvariable=self.v_sku, width=44, font=('Consolas', 10)).grid(row=3, column=0, columnspan=2, sticky='we')
        label('Physical Stock QTY', 4, 0)
        label('Committed Stock', 4, 1)
        ttk.Spinbox(self, from_=0, to=10**9, textvariable=self.v_phys, width=18, font=FONT).grid(row=5, column=0, sticky='w', padx=(0, 8))
        ttk.Spinbox(self, from_=0, to=10**9, textvariable=self.v_comm, width=18, font=FONT).grid(row=5, column=1, sticky='w')
        ttk.Checkbutton(self, text='Active', variable=self.v_active).grid(row=6, column=0, sticky='w', pady=(10, 0))

        self.preview = tk.Label(self, bg=CARD, fg=MUTED, font=FONT, anchor='w')
        self.preview.grid(row=7, column=0, columnspan=2, sticky='we', pady=(8, 0))
        self.error = tk.Label(self, bg='#fef2f2', fg='#b91c1c', font=FONT, anchor='w', justify='left', wraplength=380)
        btns = tk.Frame(self, bg=CARD)
        btns.grid(row=9, column=0, columnspan=2, sticky='e', pady=(14, 0))
        ttk.Button(btns, text='Cancel', command=self.destroy).pack(side='right')
        ttk.Button(btns, text='Save', style='Accent.TButton', command=self.save).pack(side='right', padx=(0, 6))

        self.v_phys.trace_add('write', lambda *_: self.update_preview())
        self.v_comm.trace_add('write', lambda *_: self.update_preview())
        self.update_preview()
        self.bind('<Return>', lambda e: self.save())
        self.bind('<Escape>', lambda e: self.destroy())

        self.update_idletasks()
        x = app.winfo_rootx() + (app.winfo_width() - self.winfo_width()) // 2
        y = app.winfo_rooty() + (app.winfo_height() - self.winfo_height()) // 3
        self.geometry(f'+{max(x, 0)}+{max(y, 0)}')
        self.grab_set()
        e_name.focus_set()

    def update_preview(self):
        p, c = parse_qty(self.v_phys.get()), parse_qty(self.v_comm.get())
        if p is None or c is None or p < 0 or c < 0:
            self.preview.config(text='')
        else:
            self.preview.config(text=f'Available: {p - c:,}   ·   Status: {compute_status(p, c)}')

    def save(self):
        try:
            self.app.store.save_item(self.item['id'] if self.item else None, self.v_name.get(), self.v_sku.get(),
                                     parse_qty(self.v_phys.get()), parse_qty(self.v_comm.get()), self.v_active.get())
        except ValueError as ex:
            self.error.config(text=str(ex), padx=8, pady=6)
            self.error.grid(row=8, column=0, columnspan=2, sticky='we', pady=(8, 0))
            return
        sku = self.v_sku.get().strip()
        self.destroy()
        self.app.refresh()
        self.app.set_status(f'{"Updated" if self.item else "Added"} {sku}')


class TextDialog(tk.Toplevel):
    """Simple read-only message window with scrollable details."""

    def __init__(self, app, title, summary, lines):
        super().__init__(app)
        self.title(title)
        self.configure(bg=CARD, padx=16, pady=12)
        self.transient(app)
        tk.Label(self, text=summary, bg=CARD, font=('Segoe UI Semibold', 11), anchor='w').pack(fill='x')
        if lines:
            frame = tk.Frame(self, bg=CARD)
            frame.pack(fill='both', expand=True, pady=(8, 0))
            txt = tk.Text(frame, width=80, height=min(len(lines), 14), font=FONT, wrap='word', relief='solid', bd=1)
            sb = ttk.Scrollbar(frame, command=txt.yview)
            txt.configure(yscrollcommand=sb.set)
            txt.pack(side='left', fill='both', expand=True)
            sb.pack(side='right', fill='y')
            txt.insert('1.0', '\n'.join(lines))
            txt.configure(state='disabled')
        ttk.Button(self, text='OK', command=self.destroy).pack(anchor='e', pady=(10, 0))
        self.bind('<Escape>', lambda e: self.destroy())
        self.bind('<Return>', lambda e: self.destroy())
        self.grab_set()


class App(tk.Tk):
    def __init__(self, db_path=DB_PATH):
        super().__init__()
        self.title('Item Checker')
        self.geometry('1180x720')
        self.minsize(900, 560)
        self.configure(bg=BG)
        self.store = Store(db_path)
        self.sort_col, self.sort_dir = 'item_name', 'asc'
        self.last_result = None

        self.setup_style()
        self.build_menu()
        self.build_header()
        self.build_banner()
        self.build_body()
        self.build_statusbar()
        self.bind_keys()
        self.refresh()
        self.set_status(f'Database: {self.store.path}')
        self.protocol('WM_DELETE_WINDOW', self.on_close)

    # ---------- setup ----------
    def setup_style(self):
        s = ttk.Style(self)
        if 'vista' in s.theme_names():
            s.theme_use('vista')
        self.option_add('*Font', FONT)
        s.configure('.', font=FONT)
        s.configure('Treeview', rowheight=28, font=FONT)
        s.configure('Treeview.Heading', font=('Segoe UI Semibold', 10))
        s.configure('Accent.TButton', font=('Segoe UI Semibold', 10))
        s.configure('TNotebook.Tab', padding=(14, 5), font=('Segoe UI Semibold', 10))

    def build_menu(self):
        m = tk.Menu(self)
        f = tk.Menu(m, tearoff=0)
        f.add_command(label='Import CSV / Excel / .db backup…', accelerator='Ctrl+I', command=self.do_import)
        f.add_command(label='Download import template…', command=self.do_template)
        f.add_separator()
        f.add_command(label='Export to Excel (.xlsx)…', command=lambda: self.do_export('xlsx'))
        f.add_command(label='Export to CSV (.csv)…', command=lambda: self.do_export('csv'))
        f.add_command(label='SQLite backup (.db)…', command=lambda: self.do_export('db'))
        f.add_separator()
        f.add_command(label='Open data folder', command=lambda: os.startfile(os.path.dirname(os.path.abspath(self.store.path))))
        f.add_separator()
        f.add_command(label='Exit', command=self.on_close)
        m.add_cascade(label='File', menu=f)
        it = tk.Menu(m, tearoff=0)
        it.add_command(label='Add Item', accelerator='Ctrl+N', command=self.add_item)
        it.add_command(label='Edit Item', accelerator='Enter', command=self.edit_item)
        it.add_command(label='Toggle Active', accelerator='Space', command=self.toggle_item)
        it.add_command(label='Delete Item', accelerator='Del', command=self.delete_item)
        it.add_separator()
        it.add_command(label='Refresh', accelerator='F5', command=self.refresh)
        m.add_cascade(label='Items', menu=it)
        self.config(menu=m)
        self.row_menu = it

    def build_header(self):
        h = self.header = tk.Frame(self, bg=TEAL, padx=16, pady=10)
        h.pack(fill='x')
        left = tk.Frame(h, bg=TEAL)
        left.pack(side='left')
        tk.Label(left, text='Item Checker', bg=TEAL, fg='white', font=('Segoe UI Semibold', 16)).pack(anchor='w')
        tk.Label(left, text='Physical vs Committed stock · SQLite', bg=TEAL, fg='#ccfbf1', font=('Segoe UI', 9)).pack(anchor='w')
        ttk.Button(h, text='+ Add Item', style='Accent.TButton', command=self.add_item).pack(side='right', padx=(6, 0))
        exp = ttk.Menubutton(h, text='Export')
        em = tk.Menu(exp, tearoff=0)
        em.add_command(label='Excel (.xlsx)', command=lambda: self.do_export('xlsx'))
        em.add_command(label='CSV (.csv)', command=lambda: self.do_export('csv'))
        em.add_separator()
        em.add_command(label='SQLite backup (.db)', command=lambda: self.do_export('db'))
        exp['menu'] = em
        exp.pack(side='right', padx=(6, 0))
        ttk.Button(h, text='Import', command=self.do_import).pack(side='right')

    def build_banner(self):
        self.banner = tk.Frame(self, bg='#fffbeb', highlightbackground='#fcd34d', highlightthickness=1, padx=12, pady=8)
        tk.Label(self.banner, text='The app started with 5 sample items so you can try it out.', bg='#fffbeb', fg='#78350f').pack(side='left')
        ttk.Button(self.banner, text='Keep', command=self.keep_sample).pack(side='right')
        ttk.Button(self.banner, text='Clear sample data', command=self.clear_sample).pack(side='right', padx=(0, 6))

    def build_body(self):
        self.body = tk.Frame(self, bg=BG, padx=14, pady=10)
        self.body.pack(fill='both', expand=True)
        self.nb = ttk.Notebook(self.body)
        self.nb.pack(fill='both', expand=True)
        self.tab_items = tk.Frame(self.nb, bg=BG, padx=10, pady=10)
        self.tab_sql = tk.Frame(self.nb, bg=BG, padx=10, pady=10)
        self.nb.add(self.tab_items, text='Items')
        self.nb.add(self.tab_sql, text='SQL Console')
        self.build_items_tab()
        self.build_sql_tab()

    def stat_card(self, parent, col, title, color=None, on_click=None):
        card = tk.Frame(parent, bg=CARD, highlightbackground='#e2e8f0', highlightthickness=1, padx=14, pady=10)
        card.grid(row=0, column=col, sticky='nsew', padx=(0 if col == 0 else 8, 0))
        t = tk.Label(card, text=title.upper(), bg=CARD, fg=color or MUTED, font=('Segoe UI Semibold', 8))
        t.pack(anchor='w')
        v = tk.Label(card, text='0', bg=CARD, fg=color or '#0f172a', font=('Segoe UI Semibold', 18))
        v.pack(anchor='w')
        if on_click:
            for w in (card, t, v):
                w.bind('<Button-1>', lambda e: on_click())
                w.configure(cursor='hand2')
        return v

    def build_items_tab(self):
        t = self.tab_items
        cards = tk.Frame(t, bg=BG)
        cards.pack(fill='x')
        for i in range(5):
            cards.columnconfigure(i, weight=1, uniform='c')
        self.s_total = self.stat_card(cards, 0, 'Total Items', on_click=lambda: self.quick_filter('all', 'all'))
        self.s_active = self.stat_card(cards, 1, 'Active', on_click=lambda: self.quick_filter('1', 'all'))
        self.s_phys = self.stat_card(cards, 2, 'Physical QTY')
        self.s_comm = self.stat_card(cards, 3, 'Committed QTY')
        self.s_over = self.stat_card(cards, 4, 'Over-committed', RED, on_click=lambda: self.quick_filter('all', 'OVER-COMMITTED'))

        bar = tk.Frame(t, bg=BG, pady=10)
        bar.pack(fill='x')
        tk.Label(bar, text='Search', bg=BG, fg=MUTED).pack(side='left')
        self.v_search = tk.StringVar()
        self.e_search = ttk.Entry(bar, textvariable=self.v_search, width=34)
        self.e_search.pack(side='left', padx=(6, 12))
        self.v_search.trace_add('write', lambda *_: self.refresh())
        self.active_opts = {'All items': 'all', 'Active only': '1', 'Inactive only': '0'}
        self.v_active = tk.StringVar(value='All items')
        cb = ttk.Combobox(bar, textvariable=self.v_active, values=list(self.active_opts), state='readonly', width=14)
        cb.pack(side='left')
        cb.bind('<<ComboboxSelected>>', lambda e: self.refresh())
        self.v_statusf = tk.StringVar(value='All statuses')
        cb2 = ttk.Combobox(bar, textvariable=self.v_statusf, values=['All statuses'] + STATUSES, state='readonly', width=18)
        cb2.pack(side='left', padx=(8, 0))
        cb2.bind('<<ComboboxSelected>>', lambda e: self.refresh())
        ttk.Button(bar, text='Reset', command=lambda: self.quick_filter('all', 'all')).pack(side='left', padx=(8, 0))
        self.l_count = tk.Label(bar, bg=BG, fg=MUTED)
        self.l_count.pack(side='right')

        table = tk.Frame(t, bg=BG)
        table.pack(fill='both', expand=True)
        self.cols = [('item_name', 'Item Name', 260, 'w'), ('sku', 'SKU', 130, 'w'),
                     ('physical_stock', 'Physical', 90, 'e'), ('committed_stock', 'Committed', 95, 'e'),
                     ('available_stock', 'Available', 90, 'e'), ('status', 'Status', 140, 'w'),
                     ('active', 'Active', 70, 'center'), ('updated_at', 'Updated', 150, 'w')]
        self.tree = ttk.Treeview(table, columns=[c[0] for c in self.cols], show='headings', selectmode='browse')
        for key, title, width, anchor in self.cols:
            self.tree.heading(key, text=title, anchor=anchor, command=lambda k=key: self.sort_by(k))
            self.tree.column(key, width=width, anchor=anchor, stretch=key == 'item_name')
        for st, color in ROW_COLORS.items():
            self.tree.tag_configure(st, background=color)
        self.tree.tag_configure('inactive', foreground='#94a3b8')
        sb = ttk.Scrollbar(table, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        sb.pack(side='right', fill='y')
        self.tree.bind('<Double-1>', lambda e: self.edit_item() if self.tree.identify_row(e.y) else None)
        self.tree.bind('<Return>', lambda e: self.edit_item())
        self.tree.bind('<space>', lambda e: self.toggle_item())
        self.tree.bind('<Delete>', lambda e: self.delete_item())
        self.tree.bind('<Button-3>', self.on_right_click)
        self.tree.bind('<<TreeviewSelect>>', lambda e: self.update_buttons())

        self.empty = tk.Label(self.tree, bg=CARD, fg=MUTED)

        actions = tk.Frame(t, bg=BG, pady=8)
        actions.pack(fill='x')
        self.b_edit = ttk.Button(actions, text='Edit', command=self.edit_item)
        self.b_toggle = ttk.Button(actions, text='Toggle Active', command=self.toggle_item)
        self.b_delete = ttk.Button(actions, text='Delete', command=self.delete_item)
        for b in (self.b_edit, self.b_toggle, self.b_delete):
            b.pack(side='left', padx=(0, 6))
        tk.Label(actions, text='Tip: double-click a row to edit · Space toggles Active · Del deletes',
                 bg=BG, fg=MUTED, font=('Segoe UI', 9)).pack(side='left', padx=(10, 0))
        legend = tk.Frame(actions, bg=BG)
        legend.pack(side='right')
        for st in STATUSES:
            tk.Label(legend, text=f' {st} ', bg=ROW_COLORS[st], fg='#334155', font=('Segoe UI', 8),
                     highlightbackground='#cbd5e1', highlightthickness=1).pack(side='left', padx=2)

    def build_sql_tab(self):
        t = self.tab_sql
        presets = tk.Frame(t, bg=BG)
        presets.pack(fill='x')
        for label, sql in PRESETS:
            ttk.Button(presets, text=label, command=lambda s=sql: self.run_preset(s)).pack(side='left', padx=(0, 6))
        box = tk.Frame(t, bg=BG, pady=8)
        box.pack(fill='x')
        self.sql_text = tk.Text(box, height=8, font=('Consolas', 11), wrap='none', undo=True,
                                relief='solid', bd=1, padx=8, pady=6)
        self.sql_text.pack(fill='x')
        self.sql_text.insert('1.0', "SELECT * FROM v_item_check WHERE status = 'LOW';")
        self.sql_text.bind('<Control-Return>', lambda e: (self.run_sql(), 'break')[1])
        row = tk.Frame(t, bg=BG)
        row.pack(fill='x')
        ttk.Button(row, text='Run  (Ctrl+Enter)', style='Accent.TButton', command=self.run_sql).pack(side='left')
        self.b_sql_csv = ttk.Button(row, text='Save result as CSV…', command=self.save_sql_csv, state='disabled')
        self.b_sql_csv.pack(side='left', padx=(6, 0))
        self.l_sql = tk.Label(row, bg=BG, fg=MUTED)
        self.l_sql.pack(side='right')
        tk.Label(t, text='Table: items   ·   View: v_item_check (adds available_stock and status)   ·   Changes made here are saved to the database.',
                 bg=BG, fg=MUTED, font=('Segoe UI', 9)).pack(anchor='w', pady=(4, 6))
        self.sql_err = tk.Label(t, bg='#fef2f2', fg='#b91c1c', font=('Consolas', 10), anchor='w', justify='left', padx=8, pady=6)
        res = tk.Frame(t, bg=BG)
        res.pack(fill='both', expand=True)
        self.sql_tree = ttk.Treeview(res, show='headings')
        ys = ttk.Scrollbar(res, command=self.sql_tree.yview)
        xs = ttk.Scrollbar(res, orient='horizontal', command=self.sql_tree.xview)
        self.sql_tree.configure(yscrollcommand=ys.set, xscrollcommand=xs.set)
        self.sql_tree.grid(row=0, column=0, sticky='nsew')
        ys.grid(row=0, column=1, sticky='ns')
        xs.grid(row=1, column=0, sticky='ew')
        res.rowconfigure(0, weight=1)
        res.columnconfigure(0, weight=1)

    def build_statusbar(self):
        self.statusbar = tk.Label(self, bg='#e2e8f0', fg='#334155', anchor='w', padx=12, pady=3, font=('Segoe UI', 9))
        self.statusbar.pack(fill='x', side='bottom')

    def bind_keys(self):
        self.bind('<Control-n>', lambda e: self.add_item())
        self.bind('<Control-i>', lambda e: self.do_import())
        self.bind('<Control-f>', lambda e: (self.nb.select(self.tab_items), self.e_search.focus_set()))
        self.bind('<F5>', lambda e: self.refresh())

    # ---------- helpers ----------
    def set_status(self, msg):
        self.statusbar.config(text=f'  {msg}')

    def selected_id(self):
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def update_buttons(self):
        state = 'normal' if self.selected_id() else 'disabled'
        for b in (self.b_edit, self.b_toggle, self.b_delete):
            b.configure(state=state)

    def quick_filter(self, active, status):
        self.v_active.set({v: k for k, v in self.active_opts.items()}[active])
        self.v_statusf.set('All statuses' if status == 'all' else status)
        self.nb.select(self.tab_items)
        self.v_search.set('')  # triggers refresh

    def sort_by(self, col):
        if self.sort_col == col:
            self.sort_dir = 'desc' if self.sort_dir == 'asc' else 'asc'
        else:
            self.sort_col, self.sort_dir = col, 'asc'
        self.refresh()

    # ---------- render ----------
    def refresh(self):
        s = self.store
        t = s.totals()
        self.s_total.config(text=f'{t["total"]:,}')
        self.s_active.config(text=f'{t["active"]:,}')
        self.s_phys.config(text=f'{t["physical"]:,}')
        self.s_comm.config(text=f'{t["committed"]:,}')
        self.s_over.config(text=f'{t["over"]:,}')

        status = self.v_statusf.get()
        rows = s.list_items(self.v_search.get().strip(), self.active_opts.get(self.v_active.get(), 'all'),
                            status if status in STATUSES else 'all', self.sort_col, self.sort_dir)
        keep = self.selected_id()
        self.tree.delete(*self.tree.get_children())
        for r in rows:
            tags = [r['status']] + ([] if r['active'] else ['inactive'])
            self.tree.insert('', 'end', iid=str(r['id']), tags=tags, values=(
                r['item_name'], r['sku'], f'{r["physical_stock"]:,}', f'{r["committed_stock"]:,}',
                f'{r["available_stock"]:,}', r['status'], 'Yes' if r['active'] else 'No', r['updated_at']))
        if keep and self.tree.exists(str(keep)):
            self.tree.selection_set(str(keep))
            self.tree.see(str(keep))
        self.l_count.config(text=f'Showing {len(rows):,} of {t["total"]:,}')

        for key, title, _, anchor in self.cols:
            arrow = (' ▲' if self.sort_dir == 'asc' else ' ▼') if key == self.sort_col else ''
            self.tree.heading(key, text=title + arrow)

        if rows:
            self.empty.place_forget()
        else:
            self.empty.config(text='No items match the current filters.' if t['total']
                              else 'No items yet. Click "+ Add Item" or Import a file.')
            self.empty.place(relx=0.5, rely=0.4, anchor='center')

        if s.sample_flag:
            self.banner.pack(fill='x', padx=14, pady=(10, 0), after=self.header)
        else:
            self.banner.pack_forget()
        self.update_buttons()

    # ---------- item actions ----------
    def add_item(self):
        ItemDialog(self)

    def edit_item(self):
        item_id = self.selected_id()
        if item_id:
            item = self.store.get(item_id)
            if item:
                ItemDialog(self, item)

    def toggle_item(self):
        item_id = self.selected_id()
        if item_id:
            self.store.toggle_active(item_id)
            self.refresh()

    def delete_item(self):
        item_id = self.selected_id()
        if not item_id:
            return
        r = self.store.get(item_id)
        if r and messagebox.askyesno('Delete item?', f'{r["item_name"]} ({r["sku"]}) will be permanently removed.',
                                     icon='warning', parent=self):
            self.store.delete(item_id)
            self.refresh()
            self.set_status(f'Deleted {r["sku"]}')

    def on_right_click(self, e):
        row = self.tree.identify_row(e.y)
        if row:
            self.tree.selection_set(row)
            self.row_menu.tk_popup(e.x_root, e.y_root)

    def clear_sample(self):
        if messagebox.askyesno('Clear sample data?', 'All items currently in the list will be removed so you can start fresh.',
                               icon='warning', parent=self):
            self.store.clear_all()
            self.refresh()
            self.set_status('Sample data cleared')

    def keep_sample(self):
        self.store.sample_flag = False
        self.refresh()

    # ---------- import / export ----------
    def do_import(self):
        path = filedialog.askopenfilename(parent=self, title='Import items', filetypes=[
            ('Item files', '*.xlsx *.xlsm *.csv *.db *.sqlite *.sqlite3'), ('Excel', '*.xlsx *.xlsm'), ('CSV', '*.csv'),
            ('SQLite backup', '*.db *.sqlite *.sqlite3'), ('All files', '*.*')])
        if not path:
            return
        try:
            if os.path.splitext(path)[1].lower() in ('.db', '.sqlite', '.sqlite3'):
                if not messagebox.askyesno('Restore backup?', 'Restoring a .db backup REPLACES ALL current items.\n\nContinue?',
                                           icon='warning', parent=self):
                    return
                n = self.store.restore_from(path)
                self.refresh()
                self.set_status(f'Backup restored from {os.path.basename(path)}')
                messagebox.showinfo('Backup restored', f'{n:,} items loaded.', parent=self)
                return
            added, updated, skipped = self.store.import_rows(read_table_file(path))
        except Exception as ex:
            messagebox.showerror('Import failed', str(ex), parent=self)
            return
        self.refresh()
        summary = f'{added:,} added, {updated:,} updated, {len(skipped):,} skipped.'
        self.set_status(f'Imported {os.path.basename(path)}: {summary}')
        lines = [f'Row {n}{f" ({sku})" if sku else ""}: {why}' for n, sku, why in skipped]
        TextDialog(self, 'Import finished', summary, lines)

    def ask_save(self, title, ext, label):
        return filedialog.asksaveasfilename(parent=self, title=title, defaultextension=ext,
                                            initialfile=f'ItemChecker_{today()}{ext}', filetypes=[(label, f'*{ext}')])

    def do_export(self, kind):
        try:
            if kind == 'db':
                path = self.ask_save('Save SQLite backup', '.db', 'SQLite database')
                if path:
                    self.store.backup_to(path)
            else:
                if kind == 'xlsx' and openpyxl is None:
                    raise RuntimeError('Excel export needs openpyxl.  Run:  pip install openpyxl')
                path = self.ask_save('Export items', f'.{kind}', 'Excel workbook' if kind == 'xlsx' else 'CSV file')
                if path:
                    rows = self.store.export_rows()
                    if kind == 'xlsx':
                        write_xlsx(path, rows)
                    else:
                        write_csv(path, HEADERS, [(r['item_name'], r['sku'], r['physical_stock'], r['committed_stock'],
                                                   'TRUE' if r['active'] else 'FALSE') for r in rows])
        except Exception as ex:
            messagebox.showerror('Export failed', str(ex), parent=self)
            return
        if path:
            self.set_status(f'Saved {path}')

    def do_template(self):
        if openpyxl is None:
            messagebox.showerror('Template', 'Excel support needs openpyxl.  Run:  pip install openpyxl', parent=self)
            return
        path = filedialog.asksaveasfilename(parent=self, title='Save import template', defaultextension='.xlsx',
                                            initialfile='ItemChecker_Import_Template.xlsx', filetypes=[('Excel workbook', '*.xlsx')])
        if path:
            try:
                write_template(path)
                self.set_status(f'Template saved: {path}')
            except Exception as ex:
                messagebox.showerror('Template', str(ex), parent=self)

    # ---------- SQL console ----------
    def run_preset(self, sql):
        self.sql_text.delete('1.0', 'end')
        self.sql_text.insert('1.0', sql)
        self.run_sql()

    def run_sql(self):
        sql = self.sql_text.get('1.0', 'end').strip()
        if not sql:
            return
        self.sql_err.pack_forget()
        started = datetime.datetime.now()
        try:
            cols, rows, changed, any_change = self.store.run_sql(sql)
        except sqlite3.Error as ex:
            if self.store.conn.in_transaction:
                self.store.conn.execute('ROLLBACK')
            self.sql_err.config(text=f'Error: {ex}')
            self.sql_err.pack(fill='x', before=self.sql_tree.master, pady=(0, 6))
            self.show_sql_result(None, None)
            self.l_sql.config(text='')
            return
        ms = (datetime.datetime.now() - started).total_seconds() * 1000
        self.show_sql_result(cols, rows)
        parts = []
        if cols is not None:
            parts.append(f'{len(rows):,} row(s)')
        if changed:
            parts.append(f'{changed:,} changed')
        elif cols is None:
            parts.append('Done')
        parts.append(f'{ms:.1f} ms')
        self.l_sql.config(text='  ·  '.join(parts))
        if any_change:
            self.refresh()

    def show_sql_result(self, cols, rows):
        tv = self.sql_tree
        tv.delete(*tv.get_children())
        self.last_result = (cols, rows) if cols is not None else None
        self.b_sql_csv.configure(state='normal' if self.last_result else 'disabled')
        if cols is None:
            tv['columns'] = ()
            return
        tv['columns'] = [f'c{i}' for i in range(len(cols))]
        for i, c in enumerate(cols):
            numeric = any(isinstance(r[i], (int, float)) for r in rows[:50]) and all(
                r[i] is None or isinstance(r[i], (int, float)) for r in rows[:50])
            tv.heading(f'c{i}', text=c, anchor='e' if numeric else 'w')
            tv.column(f'c{i}', anchor='e' if numeric else 'w', width=max(90, min(360, len(c) * 10 + 40)), stretch=True)
        for r in rows[:5000]:
            tv.insert('', 'end', values=['NULL' if v is None else (f'{v:,}' if isinstance(v, int) else v) for v in r])

    def save_sql_csv(self):
        if not self.last_result:
            return
        path = filedialog.asksaveasfilename(parent=self, title='Save query result', defaultextension='.csv',
                                            initialfile=f'ItemChecker_query_{today()}.csv', filetypes=[('CSV file', '*.csv')])
        if path:
            write_csv(path, *self.last_result)
            self.set_status(f'Saved {path}')

    def on_close(self):
        self.store.close()
        self.destroy()


def main():
    if sys.platform == 'win32':
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    try:
        app = App()
    except sqlite3.Error as ex:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror('Item Checker', f'Could not open the database:\n{DB_PATH}\n\n{ex}')
        return
    app.mainloop()


if __name__ == '__main__':
    main()
